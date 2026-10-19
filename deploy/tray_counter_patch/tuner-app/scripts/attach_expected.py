"""Attach expected results (e.g. expected_result_tray_v4.xlsx) to the tuner's runs, so side by side can show the
expected running count between existing and patched. Not ground truth: the renderer's own list of crossings.

Usage: python scripts/attach_expected.py <expected.xlsx>      (run from tuner-app/; writes runs/<id>/expected.json)
A run gets the expected list of its case number, taken from its original file name (kasusNN_...) or its name (case NN)."""
import json, os, re, sys
import openpyxl

wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
summary = {r[0]: r for r in list(wb["Ringkasan"].iter_rows(values_only=True))[1:] if r[0]}
cases = {}
for r in list(wb["Detail per tray"].iter_rows(values_only=True))[1:]:
    for d, t in re.findall(r"(OUT|IN)@([\d.]+)", r[11] or ""):
        cases.setdefault(r[0], []).append(dict(t=float(t), dir=d.lower(), viol=bool(r[10]) and d == "OUT",
                                               fill=r[3], role=r[4], lane=r[5], tray=r[2]))
RUNS = "runs"
for rid in sorted(os.listdir(RUNS)):
    meta = os.path.join(RUNS, rid, "meta.json")
    if not os.path.exists(meta):
        continue
    m = json.load(open(meta))
    num = re.match(r"kasus(\d+)", m.get("originalName") or "") or re.match(r"case\s*(\d+)", m.get("name") or "", re.I)
    if not num or int(num.group(1)) not in summary:
        continue
    n = int(num.group(1)); s = summary[n]
    out = dict(source=os.path.basename(sys.argv[1]), case=n, name=s[1], out=s[8], inn=s[9], viol=s[10],
               crossings=sorted(cases.get(n, []), key=lambda c: c["t"]))
    json.dump(out, open(os.path.join(RUNS, rid, "expected.json"), "w"), indent=1)
    print(f"{m['name']:10s} <- case {n:2d} {s[1]}: {len(out['crossings'])} crossings")
