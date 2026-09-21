"""Metodologi slide for the kick-off deck, as EDITABLE PowerPoint shapes (not one flat image).

Everything except the two photos is a native shape/text box, so it can be copied
straight into the original deck. Styling follows the deck itself: Poppins for
headings, Nunito for body, #073763 / #475569 / #E2E8F0 / #F0F4F8, small type.

Occupies the LEFT half of a 16:9 slide (10 x 5.625 in); pass --side right to
place it on the other half.

    python scripts/123_methodology_slide.py [--person 511] [--side left]
Writes artifacts/report/metodologi_slide.pptx and the two PNGs it places.
"""
import argparse, csv, json, sys
from pathlib import Path
import cv2, numpy as np, torch, torchvision
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
ap = argparse.ArgumentParser()
ap.add_argument("--person", default="511")
ap.add_argument("--side", choices=["left", "right"], default="left")
ap.add_argument("--out", default="artifacts/report/metodologi_slide.pptx")
a = ap.parse_args()
OUTDIR = (ROOT/a.out).parent; OUTDIR.mkdir(parents=True, exist_ok=True)

NAVY, SLATE, BORDER, TINT = RGBColor(0x07, 0x37, 0x63), RGBColor(0x47, 0x55, 0x69), RGBColor(0xE2, 0xE8, 0xF0), RGBColor(0xF0, 0xF4, 0xF8)
OKC, BADC = RGBColor(0x1B, 0x7F, 0x4B), RGBColor(0xB4, 0x23, 0x1E)
HEAD, BODY = "Poppins", "Nunito"

# ---------------------------------------------------------------- real data + model
rows = {r["n"]: r for r in csv.DictReader(open(ROOT/"data/label_sppg/labels.csv"))}
rec = rows[a.person]
frame = cv2.imread(str(ROOT/"data/sppg"/rec["frame"])); H, W = frame.shape[:2]
box = [int(v) for v in json.loads(rec["box"])]

from ultralytics import YOLO
res = YOLO(str(ROOT/"yolo11s-pose.pt")).predict(frame, imgsz=640, conf=0.40, verbose=False)[0]
boxes = res.boxes.xyxy.cpu().numpy().astype(int) if len(res.boxes) else np.zeros((0, 4), int)

meta = json.loads((ROOT/"runs/person_ppe/export/apd_ours_person_ppe.json").read_text())
net = torchvision.models.efficientnet_b0(weights=None)
net.classifier[-1] = torch.nn.Linear(net.classifier[-1].in_features, 3)
net.load_state_dict(torch.load(ROOT/"runs/person_ppe/efficientnet_b0_320_sppg/best.pt", map_location="cpu"))
net.eval()

pw, ph = box[2]-box[0], box[3]-box[1]
crop = frame[int(max(0, box[1]-.18*ph)):int(min(H, box[3]+.18*ph)), int(max(0, box[0]-.18*pw)):int(min(W, box[2]+.18*pw))]
s = 320/max(crop.shape[:2]); rs = cv2.resize(crop, (round(crop.shape[1]*s), round(crop.shape[0]*s)))
lb = np.full((320, 320, 3), 114, np.uint8)
y0, x0 = (320-rs.shape[0])//2, (320-rs.shape[1])//2
lb[y0:y0+rs.shape[0], x0:x0+rs.shape[1]] = rs
MEAN, STD = np.array([0.485, 0.456, 0.406], np.float32), np.array([0.229, 0.224, 0.225], np.float32)
x = ((cv2.cvtColor(lb, cv2.COLOR_BGR2RGB).astype(np.float32)/255 - MEAN)/STD).transpose(2, 0, 1)
with torch.no_grad(): prob = torch.sigmoid(net(torch.from_numpy(x)[None]))[0].numpy()
TH = [meta["thresholds"][k] for k in ("head_cap", "mask", "gloves")]

vis = frame.copy()
for b in boxes:
    hit = abs(b[0]-box[0]) < 40 and abs(b[1]-box[1]) < 40
    cv2.rectangle(vis, (b[0], b[1]), (b[2], b[3]), (0, 200, 255) if hit else (235, 190, 120), 7 if hit else 5)
p_frame, p_crop = OUTDIR/"metodologi_frame.png", OUTDIR/"metodologi_crop.png"
cv2.imwrite(str(p_frame), cv2.resize(vis, (1280, round(1280*H/W))))
cv2.imwrite(str(p_crop), cv2.resize(lb, (480, 480), interpolation=cv2.INTER_CUBIC))
print({k: round(float(v), 3) for k, v in zip(("head_cap", "mask", "gloves"), prob)}, "thresholds", TH,
      f"| {len(boxes)} people detected")

# ---------------------------------------------------------------- slide
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
sl = prs.slides.add_slide(prs.slide_layouts[6])
X0 = Inches(0.42) if a.side == "left" else Inches(5.08)
CW = Inches(4.5)


def text(x, y, w, h, runs, size=8, color=SLATE, font=BODY, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = sl.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]; p.alignment = align
    for i, (t, opt) in enumerate(runs if isinstance(runs, list) else [(runs, {})]):
        r = p.add_run(); r.text = t
        f = r.font
        f.name = opt.get("font", font); f.size = Pt(opt.get("size", size)); f.bold = opt.get("bold", bold)
        f.color.rgb = opt.get("color", color)
    return tb


def step(y, n, title, sub):
    d = Inches(0.2)
    c = sl.shapes.add_shape(MSO_SHAPE.OVAL, X0, y, d, d)
    c.fill.solid(); c.fill.fore_color.rgb = NAVY; c.line.fill.background(); c.shadow.inherit = False
    tf = c.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = str(n); r.font.size = Pt(7); r.font.bold = True
    r.font.name = HEAD; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    text(X0 + Inches(0.27), y - Inches(0.015), CW, Inches(0.2),
         [(title, {"size": 10, "bold": True, "color": NAVY, "font": HEAD})])
    text(X0 + Inches(0.27), y + Inches(0.155), CW, Inches(0.16), [(sub, {"size": 7.5, "color": SLATE})])
    return y + Inches(0.4)


def arrow(y):
    ar = sl.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, X0 + Inches(0.06), y, Inches(0.08), Inches(0.13))
    ar.fill.solid(); ar.fill.fore_color.rgb = BORDER; ar.line.fill.background(); ar.shadow.inherit = False
    return y + Inches(0.2)


y = Inches(0.66)
y = step(y, 1, "Deteksi orang", "YOLO11s-pose  •  akselerator Metis (AIPU)")
iw = Inches(2.62); ih = Emu(int(iw * H / W))
sl.shapes.add_picture(str(p_frame), X0, y, width=iw)
text(X0 + iw + Inches(0.1), y + Inches(0.03), CW - iw - Inches(0.1), ih,
     [(f"{len(boxes)} orang terdeteksi\nper frame (10 fps).\n\nKotak kuning:\ncontoh orang\nyang dinilai.", {"size": 7})])
y = arrow(y + ih + Inches(0.05))

y = step(y, 2, "Potong area orang", "kotak orang + 18%  →  320 × 320 piksel")
cs = Inches(0.62)
sl.shapes.add_picture(str(p_crop), X0, y, width=cs, height=cs)
text(X0 + cs + Inches(0.1), y, CW - cs - Inches(0.1), cs,
     [("Satu crop = satu orang. ", {"size": 7, "bold": True, "color": NAVY}),
      ("Tanpa ROI kepala atau tangan: model menilai orang seutuhnya, tidak bergantung titik sendi yang meleset.", {"size": 7})],
     anchor=MSO_ANCHOR.MIDDLE)
y = arrow(y + cs + Inches(0.05))

y = step(y, 3, "Klasifikasi APD", "EfficientNet-B0 (CPU)  •  1 inferensi = 3 keluaran  •  ~35 ms/orang  •  15 MB")
rowh, gap = Inches(0.22), Inches(0.04)
for i, (nm, pr, th) in enumerate(zip(["Tutup kepala", "Masker", "Sarung tangan"], prob, TH)):
    ok = float(pr) >= th
    ry = y + i*(rowh + gap)
    card = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, X0, ry, CW, rowh)
    card.fill.solid(); card.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    card.line.color.rgb = BORDER; card.line.width = Pt(0.75); card.shadow.inherit = False
    card.adjustments[0] = 0.22
    text(X0 + Inches(0.09), ry, Inches(1.0), rowh, [(nm, {"size": 7.5, "bold": True, "color": NAVY})],
         anchor=MSO_ANCHOR.MIDDLE)
    gx, gw, gh = X0 + Inches(1.1), Inches(1.45), Inches(0.09)
    gy = ry + (rowh - gh)//2
    track = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, gx, gy, gw, gh)
    track.fill.solid(); track.fill.fore_color.rgb = BORDER; track.line.fill.background(); track.shadow.inherit = False
    fl = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, gx, gy, max(Inches(0.05), Emu(int(gw*float(pr)))), gh)
    fl.fill.solid(); fl.fill.fore_color.rgb = OKC if ok else BADC; fl.line.fill.background(); fl.shadow.inherit = False
    tick = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, gx + Emu(int(gw*th)), gy - Inches(0.025), Pt(1.2), gh + Inches(0.05))
    tick.fill.solid(); tick.fill.fore_color.rgb = NAVY; tick.line.fill.background(); tick.shadow.inherit = False
    text(gx + gw + Inches(0.05), ry, Inches(0.4), rowh,
         [(f"{float(pr):.2f}", {"size": 7.5, "bold": True, "color": NAVY})], anchor=MSO_ANCHOR.MIDDLE)
    text(X0 + CW - Inches(1.0), ry, Inches(0.92), rowh,
         [("SESUAI" if ok else "PELANGGARAN", {"size": 7, "bold": True, "color": OKC if ok else BADC})],
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
y = y + 3*(rowh + gap)
text(X0, y + Inches(0.02), CW, Inches(0.3),
     [("Ambang dari data validasi; pelanggaran dikonfirmasi lewat voting 5 frame.", {"size": 6.5})])

prs.save(ROOT/a.out)
print("wrote", ROOT/a.out, "| bottom edge", round((y + Inches(0.22))/Inches(1), 2), "in of 5.625")
