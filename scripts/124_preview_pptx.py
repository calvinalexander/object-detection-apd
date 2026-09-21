"""Rough preview of a .pptx as PNG (shapes, text, pictures) - no Office needed.

    python scripts/124_preview_pptx.py artifacts/report/metodologi_slide.pptx [out.png]
Approximate: fonts substituted with DejaVu, no auto-fit; good enough to catch
overflow and overlap.
"""
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu
from PIL import Image, ImageDraw, ImageFont

src = Path(sys.argv[1]); out = Path(sys.argv[2] if len(sys.argv) > 2 else str(src.with_suffix(".preview.png")))
prs = Presentation(str(src)); S = 160                      # px per inch
W = int(prs.slide_width/914400*S); H = int(prs.slide_height/914400*S)
F = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"; FB = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
img = Image.new("RGB", (W, H), "white"); d = ImageDraw.Draw(img)
d.rectangle([0, 0, W-1, H-1], outline=(210, 210, 210))
px = lambda v: int(Emu(v)/914400*S)


def rgb(c):
    try: return tuple(c.rgb) if c and c.type is not None else None
    except Exception: return None


for sh in prs.slides[0].shapes:
    x, y, w, h = px(sh.left), px(sh.top), px(sh.width), px(sh.height)
    if sh.shape_type == 13:                                # picture
        im = Image.open(sh.image.blob and __import__("io").BytesIO(sh.image.blob)).convert("RGB").resize((max(1, w), max(1, h)))
        img.paste(im, (x, y)); continue
    fill = rgb(sh.fill.fore_color) if sh.fill.type == 1 else None
    line = rgb(sh.line.color) if sh.line.fill.type == 1 else None
    if fill or line:
        shape = "ellipse" if "OVAL" in str(sh.shape_type) else "rect"
        (d.ellipse if shape == "ellipse" else d.rounded_rectangle)(
            [x, y, x+w, y+h], **({"radius": min(10, h//2)} if shape == "rect" else {}), fill=fill, outline=line)
    if not sh.has_text_frame: continue
    tf = sh.text_frame
    for p in tf.paragraphs:
        runs = [r for r in p.runs if r.text]
        if not runs: continue
        size = max(6, int((runs[0].font.size.pt if runs[0].font.size else 10) * S/72))
        ty = y + (h - size)//2 if str(tf.vertical_anchor).startswith("MIDDLE") or "OVAL" in str(sh.shape_type) else y
        tx = x
        align = str(p.alignment)
        total = sum(d.textlength(r.text.split("\n")[0], font=ImageFont.truetype(FB if r.font.bold else F,
                    max(6, int((r.font.size.pt if r.font.size else 10)*S/72)))) for r in runs)
        if align.startswith("CENTER"): tx = x + (w - total)//2
        if align.startswith("RIGHT"): tx = x + w - total
        for r in runs:
            fs = max(6, int((r.font.size.pt if r.font.size else 10)*S/72))
            fnt = ImageFont.truetype(FB if r.font.bold else F, fs)
            col = rgb(r.font.color) or (60, 60, 60)
            for j, ln in enumerate(r.text.split("\n")):
                d.text((tx, ty + j*int(fs*1.25)), ln, font=fnt, fill=col)
            tx += d.textlength(r.text.split("\n")[0], font=fnt)
img.save(out); print("wrote", out, img.size)
