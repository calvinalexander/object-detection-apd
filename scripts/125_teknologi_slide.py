"""Teknologi slide for the kick-off deck, as EDITABLE PowerPoint shapes.

Scope: only the two models that make a verdict -- the person detector and the
PPE classifier. Nothing about deployment, infrastructure or data preparation.
Same styling as the Metodologi slide (Poppins / Nunito, #073763 / #475569).

    python scripts/125_teknologi_slide.py [--side left]
Writes artifacts/report/teknologi_slide.pptx
"""
import argparse
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
ap = argparse.ArgumentParser()
ap.add_argument("--side", choices=["left", "right"], default="left")
ap.add_argument("--out", default="artifacts/report/teknologi_slide.pptx")
a = ap.parse_args()

NAVY, SLATE, BORDER, TINT = RGBColor(0x07, 0x37, 0x63), RGBColor(0x47, 0x55, 0x69), RGBColor(0xE2, 0xE8, 0xF0), RGBColor(0xF0, 0xF4, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEAD, BODY = "Poppins", "Nunito"

CARDS = [
    ("Deteksi orang", "YOLO11s-pose", [
        ("Arsitektur", "YOLO11s-pose (Ultralytics), 1 kelas: orang"),
        ("Masukan", "frame CCTV 640 × 640"),
        ("Keluaran", "kotak orang + 17 titik sendi (COCO-17)"),
        ("Ukuran model", "9,9 juta parameter  •  23,3 GFLOPs  •  38 MB (ONNX)"),
        ("Presisi", "INT8, dijalankan di akselerator Metis (AIPU)"),
    ]),
    ("Klasifikasi APD", "EfficientNet-B0", [
        ("Arsitektur", "EfficientNet-B0 (PyTorch/torchvision), multi-label"),
        ("Masukan", "crop 1 orang, 320 × 320 (kotak + 18%)"),
        ("Keluaran", "3 probabilitas: tutup kepala, masker, sarung tangan"),
        ("Ukuran model", "4,0 juta parameter  •  1,6 GFLOPs  •  15 MB"),
        ("Presisi", "FP32 di CPU, ~35 ms per orang (ONNX Runtime)"),
    ]),
]

prs = Presentation(); prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
sl = prs.slides.add_slide(prs.slide_layouts[6])
X0 = Inches(0.42) if a.side == "left" else Inches(5.08)
CW = Inches(4.5)


def text(x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = sl.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]; p.alignment = align
    for t, o in runs:
        r = p.add_run(); r.text = t
        r.font.name = o.get("font", BODY); r.font.size = Pt(o.get("size", 7.5))
        r.font.bold = o.get("bold", False); r.font.color.rgb = o.get("color", SLATE)
    return tb


y = Inches(0.72)
for ci, (title, model, specs) in enumerate(CARDS):
    ch = Inches(0.44) + Inches(0.205)*len(specs) + Inches(0.12)
    card = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, X0, y, CW, ch)
    card.fill.solid(); card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = BORDER; card.line.width = Pt(0.75); card.shadow.inherit = False
    card.adjustments[0] = 0.06

    # header: step number, title, model chip
    d = Inches(0.19)
    dot = sl.shapes.add_shape(MSO_SHAPE.OVAL, X0 + Inches(0.12), y + Inches(0.13), d, d)
    dot.fill.solid(); dot.fill.fore_color.rgb = NAVY; dot.line.fill.background(); dot.shadow.inherit = False
    tf = dot.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = str(ci+1); r.font.size = Pt(7); r.font.bold = True
    r.font.name = HEAD; r.font.color.rgb = WHITE
    text(X0 + Inches(0.38), y + Inches(0.11), Inches(1.6), Inches(0.24),
         [(title, {"size": 10, "bold": True, "color": NAVY, "font": HEAD})])
    chip_w = Inches(1.35)
    chip = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, X0 + CW - chip_w - Inches(0.12), y + Inches(0.1), chip_w, Inches(0.24))
    chip.fill.solid(); chip.fill.fore_color.rgb = TINT
    chip.line.color.rgb = BORDER; chip.line.width = Pt(0.75); chip.shadow.inherit = False
    chip.adjustments[0] = 0.3
    tfc = chip.text_frame; tfc.margin_left = tfc.margin_right = 0
    pc = tfc.paragraphs[0]; pc.alignment = PP_ALIGN.CENTER
    rc = pc.add_run(); rc.text = model; rc.font.size = Pt(8); rc.font.bold = True
    rc.font.name = HEAD; rc.font.color.rgb = NAVY

    sy = y + Inches(0.44)
    for k, (label, val) in enumerate(specs):
        ry = sy + Inches(0.205)*k
        text(X0 + Inches(0.18), ry, Inches(0.85), Inches(0.19),
             [(label, {"size": 7.5, "bold": True, "color": NAVY})], anchor=MSO_ANCHOR.MIDDLE)
        text(X0 + Inches(1.02), ry, CW - Inches(1.2), Inches(0.19),
             [(val, {"size": 7.5})], anchor=MSO_ANCHOR.MIDDLE)
    y = y + ch + Inches(0.22)

# the chain, one line
text(X0, y, CW, Inches(0.2),
     [("Alur: ", {"size": 7.5, "bold": True, "color": NAVY}),
      ("frame → YOLO11s-pose (AIPU) → crop orang → EfficientNet-B0 (CPU) → 3 keputusan APD", {"size": 7.5})])

prs.save(ROOT/a.out)
print("wrote", ROOT/a.out, "| bottom edge", round((y + Inches(0.2))/Inches(1), 2), "in of 5.625")
