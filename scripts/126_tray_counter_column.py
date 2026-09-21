"""Adds the Tray Counter column to the Metodologi and Teknologi slides.

Both slides currently occupy only the LEFT half of the 16:9 slide (the APD
pipeline). This script opens them, fills the RIGHT half with the tray-counter
pipeline in exactly the same visual language (Poppins / Nunito, #073763 /
#475569, numbered steps + arrows on Metodologi, spec cards on Teknologi), adds
a small section label above each column so the two pipelines can be told apart,
and saves NEW files -- the originals are left untouched.

Content is taken from artifacts/report/docs.md (sections 3, 4.1, 5.4, 6.1, 6.3,
7.1, 7.4, 9.2, 10.3, 10.4).

    python scripts/126_tray_counter_column.py
Writes artifacts/report/metodologi_slide_v2.pptx and teknologi_slide_v2.pptx
"""
import argparse
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
ap = argparse.ArgumentParser()
ap.add_argument("--metodologi-in", default="artifacts/report/metodologi_slide.pptx")
ap.add_argument("--teknologi-in", default="artifacts/report/teknologi_slide.pptx")
ap.add_argument("--metodologi-out", default="artifacts/report/metodologi_slide_v2.pptx")
ap.add_argument("--teknologi-out", default="artifacts/report/teknologi_slide_v2.pptx")
a = ap.parse_args()

NAVY, SLATE, BORDER, TINT = RGBColor(0x07, 0x37, 0x63), RGBColor(0x47, 0x55, 0x69), RGBColor(0xE2, 0xE8, 0xF0), RGBColor(0xF0, 0xF4, 0xF8)
OKC, BADC = RGBColor(0x1B, 0x7F, 0x4B), RGBColor(0xB4, 0x23, 0x1E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEAD, BODY = "Poppins", "Nunito"

LEFT_X, RIGHT_X, CW = Inches(0.42), Inches(5.08), Inches(4.5)


# ---------------------------------------------------------------- shape helpers
def text(sl, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None):
    tb = sl.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]; p.alignment = align
    if spacing: p.line_spacing = spacing
    for t, o in runs:
        r = p.add_run(); r.text = t
        r.font.name = o.get("font", BODY); r.font.size = Pt(o.get("size", 7))
        r.font.bold = o.get("bold", False); r.font.color.rgb = o.get("color", SLATE)
    return tb


def box(sl, x, y, w, h, fill=WHITE, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=None, rot=None):
    s = sl.shapes.add_shape(shape, x, y, w, h)
    if fill is None: s.fill.background()
    else: s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else: s.line.color.rgb = line; s.line.width = Pt(0.75)
    s.shadow.inherit = False
    if adj is not None: s.adjustments[0] = adj
    if rot is not None: s.rotation = rot
    return s


def eyebrow(sl, x, label, note):
    """Small section label + hairline rule above a column."""
    text(sl, x, Inches(0.31), CW, Inches(0.16),
         [(label, {"size": 8, "bold": True, "color": NAVY, "font": HEAD}),
          ("   " + note, {"size": 8, "color": SLATE, "font": HEAD})])
    box(sl, x, Inches(0.5), CW, Pt(0.9), fill=BORDER, line=None, shape=MSO_SHAPE.RECTANGLE)


def step(sl, x, y, n, title, sub):
    d = Inches(0.2)
    c = box(sl, x, y, d, d, fill=NAVY, line=None, shape=MSO_SHAPE.OVAL)
    tf = c.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = str(n); r.font.size = Pt(7); r.font.bold = True
    r.font.name = HEAD; r.font.color.rgb = WHITE
    text(sl, x + Inches(0.27), y - Inches(0.015), CW, Inches(0.2),
         [(title, {"size": 10, "bold": True, "color": NAVY, "font": HEAD})])
    text(sl, x + Inches(0.27), y + Inches(0.155), CW, Inches(0.16), [(sub, {"size": 7.5})])
    return y + Inches(0.4)


def arrow(sl, x, y):
    box(sl, x + Inches(0.06), y, Inches(0.08), Inches(0.12), fill=BORDER, line=None, shape=MSO_SHAPE.DOWN_ARROW)
    return y + Inches(0.17)


# ---------------------------------------------------------------- metodologi: right column
def build_metodologi(sl):
    X = RIGHT_X
    eyebrow(sl, LEFT_X, "PEMANTAUAN APD", "—  per orang, dari stream langsung")
    eyebrow(sl, X, "PENGHITUNG NAMPAN", "—  per nampan, dari rekaman")

    y = Inches(0.65)

    y = step(sl, X, y, 1, "Rekam per segmen", "CameraRecorder  •  1280 × 720, 15 fps  •  rotasi 2 menit")
    text(sl, X, y, CW, Inches(0.26),
         [("Nampan dihitung dari file .mp4 lokal, bukan dari stream: tanpa frame hilang, urutan frame utuh, sehingga tidak ada nampan yang “melompati” garis.", {})])
    y = arrow(sl, X, y + Inches(0.26))

    y = step(sl, X, y, 2, "Deteksi nampan", "YOLO11s 640 × 640  •  akselerator Metis (AIPU)")
    text(sl, X, y, CW, Inches(0.26),
         [("Model kelas nampan (penuh / kosong). Ambang keyakinan 0,1; kotak di bawah 2.000 piksel² dibuang agar pantulan dan remah tidak ikut terhitung.", {})])
    y = arrow(sl, X, y + Inches(0.26))

    y = step(sl, X, y, 3, "Lacak antar-frame", "AdvancedObjectTracker  •  satu ID per nampan")
    fh = Inches(0.24)
    box(sl, X, y, CW, fh, adj=0.22)
    text(sl, X + Inches(0.09), y, CW - Inches(0.18), fh,
         [("Skor pencocokan = ", {"size": 7, "bold": True, "color": NAVY}),
          ("0,5 × IoU  +  0,25 × jarak  +  0,25 × prediksi kecepatan", {"size": 7})],
         anchor=MSO_ANCHOR.MIDDLE)
    y = arrow(sl, X, y + fh)

    y = step(sl, X, y, 4, "Hitung lintas garis", "SmartLineCrossingCounter  •  garis dari dashboard (ROI)")
    y = crossing_diagram(sl, X, y)
    y = arrow(sl, X, y)

    y = step(sl, X, y, 5, "Laporan & sinkronisasi", "SQLite  →  API  +  S3/MinIO")
    text(sl, X, y, CW, Inches(0.26),
         [("Rekap per menit dan total masuk/keluar per sesi dikirim ke API tray-counter; gambar dan klip pelanggaran diunggah ke S3.", {})])
    y += Inches(0.26)
    text(sl, X, y + Inches(0.03), CW, Inches(0.3),
         [("Garis hitung disegarkan dari API tiap 30 frame; video _out.mp4 berisi overlay hitungan disimpan sebagai bukti audit.", {"size": 6.5})])
    return y + Inches(0.22)


def crossing_diagram(sl, X, y):
    """Schematic of trays crossing the counting line -- native shapes, editable."""
    bh = Inches(0.8)
    box(sl, X, y, CW, bh, fill=TINT, adj=0.08)
    ix = lambda v: X + Inches(v)          # inches inside the box
    iy = lambda v: y + Inches(v)

    # tray glyphs
    for tx, ty in ((0.18, 0.228), (1.92, 0.493)):
        box(sl, ix(tx), iy(ty), Inches(0.28), Inches(0.14), adj=0.25)
    text(sl, ix(0.18), iy(0.40), Inches(0.5), Inches(0.1), [("nampan", {"size": 6})])

    # in / out arrows, each crossing the tilted line
    ar = box(sl, ix(0.54), iy(0.228), Inches(1.07), Inches(0.14), fill=OKC, line=None, shape=MSO_SHAPE.RIGHT_ARROW)
    ar.adjustments[0] = 0.3
    text(sl, ix(1.66), iy(0.238), Inches(0.6), Inches(0.12), [("MASUK", {"size": 6.5, "bold": True, "color": OKC})])
    ar = box(sl, ix(1.00), iy(0.493), Inches(0.85), Inches(0.14), fill=SLATE, line=None, shape=MSO_SHAPE.LEFT_ARROW)
    ar.adjustments[0] = 0.3
    text(sl, ix(1.00), iy(0.645), Inches(0.7), Inches(0.12), [("KELUAR", {"size": 6.5, "bold": True, "color": SLATE})])

    # counting line last, so it reads as lying across both arrows
    box(sl, ix(1.44), iy(0.07), Pt(1.6), Inches(0.66), fill=NAVY, line=None, shape=MSO_SHAPE.RECTANGLE, rot=22)
    text(sl, ix(1.58), iy(0.03), Inches(0.8), Inches(0.1), [("garis hitung", {"size": 6, "bold": True, "color": NAVY})])

    # explanation, right half of the box
    text(sl, ix(2.42), iy(0.1), Inches(1.95), Inches(0.6),
         [("Jarak bertanda", {"size": 6.5, "bold": True, "color": NAVY}),
          (" tiap nampan ke garis dihaluskan Kalman. Tanda berubah → satu penyeberangan; arah perubahan menentukan MASUK atau KELUAR. Jeda 30 frame mencegah hitungan ganda.", {"size": 6.5})],
         spacing=1.15)
    return y + bh


# ---------------------------------------------------------------- teknologi: right column
CARDS = [
    ("Deteksi nampan", "YOLO11s", [
        ("Arsitektur", "YOLO11s (Ultralytics), kelas nampan penuh / kosong"),
        ("Masukan", "frame rekaman 1280 × 720 → letterbox 640 × 640"),
        ("Keluaran", "kotak nampan + keyakinan + kelas"),
        ("Ambang", "conf 0,1  •  NMS IoU 0,45  •  area min 2.000 piksel²"),
        ("Presisi", "INT8 .axnet (Voyager SDK), akselerator Metis (AIPU)"),
    ]),
    ("Pelacakan & hitung", "IoU + Kalman", [
        ("Pelacak", "AdvancedObjectTracker (kustom), ID stabil per nampan"),
        ("Skor cocok", "0,5 × IoU + 0,25 × jarak + 0,25 × prediksi kecepatan"),
        ("Penghitung", "SmartLineCrossingCounter, jarak bertanda + Kalman1D"),
        ("Parameter", "IoU 0,2  •  jarak maks 100 px  •  anti-pantul 30 frame"),
        ("Keluaran", "event MASUK / KELUAR, rekap per menit ke SQLite"),
    ]),
]


def build_teknologi(sl):
    X = RIGHT_X
    eyebrow(sl, LEFT_X, "PEMANTAUAN APD", "—  dua model, satu keputusan per orang")
    eyebrow(sl, X, "PENGHITUNG NAMPAN", "—  satu model + pelacak kustom")

    y = Inches(0.72)
    for ci, (title, chip_t, specs) in enumerate(CARDS):
        ch = Inches(0.44) + Inches(0.205) * len(specs) + Inches(0.12)
        box(sl, X, y, CW, ch, adj=0.06)

        d = Inches(0.19)
        dot = box(sl, X + Inches(0.12), y + Inches(0.13), d, d, fill=NAVY, line=None, shape=MSO_SHAPE.OVAL)
        tf = dot.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = str(ci + 1); r.font.size = Pt(7); r.font.bold = True
        r.font.name = HEAD; r.font.color.rgb = WHITE
        text(sl, X + Inches(0.38), y + Inches(0.11), Inches(1.9), Inches(0.24),
             [(title, {"size": 10, "bold": True, "color": NAVY, "font": HEAD})])

        chip_w = Inches(1.35)
        chip = box(sl, X + CW - chip_w - Inches(0.12), y + Inches(0.1), chip_w, Inches(0.24), fill=TINT, adj=0.3)
        tfc = chip.text_frame; tfc.margin_left = tfc.margin_right = 0
        pc = tfc.paragraphs[0]; pc.alignment = PP_ALIGN.CENTER
        rc = pc.add_run(); rc.text = chip_t; rc.font.size = Pt(8); rc.font.bold = True
        rc.font.name = HEAD; rc.font.color.rgb = NAVY

        sy = y + Inches(0.44)
        for k, (label, val) in enumerate(specs):
            ry = sy + Inches(0.205) * k
            text(sl, X + Inches(0.18), ry, Inches(0.85), Inches(0.19),
                 [(label, {"size": 7.5, "bold": True, "color": NAVY})], anchor=MSO_ANCHOR.MIDDLE)
            text(sl, X + Inches(1.02), ry, CW - Inches(1.2), Inches(0.19),
                 [(val, {"size": 7.5})], anchor=MSO_ANCHOR.MIDDLE)
        y = y + ch + Inches(0.22)

    text(sl, X, y, CW, Inches(0.2),
         [("Alur: ", {"size": 7.5, "bold": True, "color": NAVY}),
          ("rekaman .mp4 → YOLO11s (AIPU) → pelacak → garis hitung → laporan per menit → API/S3", {"size": 7.5})])
    return y + Inches(0.2)


# ---------------------------------------------------------------- run
for src, dst, build in ((a.metodologi_in, a.metodologi_out, build_metodologi),
                        (a.teknologi_in, a.teknologi_out, build_teknologi)):
    prs = Presentation(str(ROOT / src))
    bottom = build(prs.slides[0])
    prs.save(ROOT / dst)
    print("wrote", ROOT / dst, "| right column bottom", round(bottom / Inches(1), 2), "in of 5.625")
