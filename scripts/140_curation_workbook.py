"""Curation workbook for artifacts/filter16, in the same shape as passed_filenames.csv.

Columns are the reviewer's own: No. | Filename | Status | Keterangan.
One sheet per folder, verdicts from the previous pass carried into Status.

Colours
  red     the reviewer marked "skip" and the model now alerts anyway
  orange  marked "flag" and the model still passes
  yellow  not checked yet

    python scripts/140_curation_workbook.py
"""
import csv, os, re
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
F = ROOT/"artifacts/filter16"
OUT = ROOT/"artifacts/_ss/passed_filenames.xlsx"
eid = lambda f: re.sub(r"\.jpg$", "", f).split("_")[-1]

prev = {}
for r in csv.DictReader(open(ROOT/"artifacts/_ss/passed_filenames.csv")):
    if r["Status"].strip() in ("skip", "flag"):
        prev[eid(r["Filename"].strip())] = (r["Status"].strip(), r["Keterangan"].strip())

fill = {k: PatternFill("solid", fgColor=v) for k, v in
        dict(red="FFC7CE", orange="FCE4D6", yellow="FFF2CC").items()}
wb = Workbook(); wb.remove(wb.active)
for folder in ("flagged", "passed"):
    ws = wb.create_sheet(folder)
    ws.append(["No.", "Filename", "Status", "Keterangan"])
    files = sorted(os.listdir(F/folder))
    n = {"red": 0, "orange": 0, "yellow": 0}
    for i, fn in enumerate(files, 1):
        status, note = prev.get(eid(fn), ("", ""))
        ws.append([i, fn, status, note])
        tag = ("yellow" if not status else
               "red" if folder == "flagged" and status == "skip" else
               "orange" if folder == "passed" and status == "flag" else None)
        if tag:
            n[tag] += 1
            for c in range(1, 5): ws.cell(i+1, c).fill = fill[tag]
    ws.append([])
    ws.append(["Total files", len(files)])
    ws.append(["Flag", sum(1 for f in files if prev.get(eid(f), ("",))[0] == "flag")])
    ws.append(["Skip", sum(1 for f in files if prev.get(eid(f), ("",))[0] == "skip")])
    ws.append(["Belum diisi", n["yellow"]])
    if folder == "flagged":
        ws.append(["Merah = dulu skip, sekarang kena alert", n["red"]])
    else:
        ws.append(["Oranye = dulu flag, masih lolos", n["orange"]])
    for c, w in zip(range(1, 5), (6, 46, 10, 30)): ws.column_dimensions[get_column_letter(c)].width = w
    for c in range(1, 5): ws.cell(1, c).font = Font(bold=True)
    ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:D{len(files)+1}"
    print(f"{folder:8} {len(files):5d} files  red {n['red']}  orange {n['orange']}  belum dicek {n['yellow']}")
wb.save(OUT)
print(OUT)
